"""Čítanie súboru na import: CSV alebo Excel na riadky s hodnotami.

Súbor píše človek v Exceli, nie program, takže sa číta zhovievavo:
hlavičky bez ohľadu na diakritiku a synonymá, sumy s čiarkou aj bodkou,
dátumy v slovenskom aj ISO tvare, stav aj po slovensky. Čo sa prečítať
nedá, nie je výnimka, ale chyba pri riadku, ktorú rozhranie ukáže.

Databázu tu nič nevidí; dohľadanie setov a zápis sú v ``importer.py``.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from lego_api.models import COLLECTION_FLAGS, ItemCondition, ItemPurpose

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_QUANTITY = 99
#: Poznámka vzorových riadkov šablóny. Taký riadok sa neimportuje.
EXAMPLE_NOTE = "príklad, zmaž tento riadok"


class ImportFileError(ValueError):
    """Súbor sa nedá prečítať vôbec (formát, veľkosť, chýba číslo setu)."""


def norm(value: Any) -> str:
    """„Kúpna cena (€)“ → „kupna_cena“. Na porovnanie hlavičiek aj hodnôt."""
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-z0-9]+", "_", text.lower())
    text = re.sub(r"_(eur|e)$", "", text.strip("_"))
    return text.strip("_")


#: Pole → mená stĺpcov, ktoré naň ukazujú (po ``norm``). Prvé meno je to,
#: ktoré píše export a šablóna.
COLUMNS: dict[str, tuple[str, ...]] = {
    "catalog_num": (
        "cislo_setu",
        "cislo",
        "set",
        "set_cislo",
        "katalogove_cislo",
        "cislo_lego",
        "lego",
        "set_number",
        "set_num",
        "catalog_num",
        "number",
        "item_number",
    ),
    "quantity": ("pocet", "pocet_kusov", "mnozstvo", "ks", "kusy", "kusov", "quantity", "qty"),
    "ownership": ("vlastnictvo", "stav_vlastnictva", "vlastnene", "status", "ownership"),
    "condition": ("stav", "stav_kusu", "condition"),
    "purpose": ("zoznam", "ucel", "purpose"),
    "location": (
        "umiestnenie",
        "ulozenie",
        "ulozene",
        "kde_ulozene",
        "kde_je",
        "miesto",
        "location",
    ),
    "box": ("krabica", "cislo_krabice", "box", "priehradka", "skatula"),
    "flags": ("priznaky", "prislusenstvo", "flags"),
    "purchase_price": (
        "kupna_cena",
        "cena",
        "nakupna_cena",
        "zaplatene",
        "cena_za_kus",
        "purchase_price",
        "price",
        "paid",
    ),
    "purchase_currency": ("mena_kupy", "mena", "currency", "purchase_currency"),
    "purchase_price_original": (
        "kupna_cena_v_mene",
        "cena_v_mene",
        "purchase_price_original",
    ),
    "purchase_date": (
        "datum_kupy",
        "datum_nakupu",
        "kupene_dna",
        "datum",
        "purchase_date",
        "date",
        "bought",
    ),
    "purchase_place": ("kde_kupene", "obchod", "kupene_v", "kupene", "purchase_place", "shop"),
    "sold_price": ("predajna_cena", "predane_za", "sold_price", "sale_price"),
    "sold_date": ("datum_predaja", "predane_dna", "sold_date"),
    "sold_via": ("kanal_predaja", "kanal", "predane_cez", "sold_via"),
    "sold_fees": ("poplatky", "sold_fees", "fees"),
    "sold_shipping": ("postovne", "sold_shipping", "shipping"),
    "target_price": ("cielova_cena", "cielova", "target_price"),
    "note": ("poznamka", "poznamky", "note", "notes"),
    "name": ("nazov", "meno", "name"),
}
_COLUMN_OF = {alias: key for key, aliases in COLUMNS.items() for alias in aliases}

OWNED, SOLD, WISH = "owned", "sold", "wish"
_OWNERSHIP = {
    OWNED: (
        "vlastnene",
        "vlastneny",
        "vlastnim",
        "mam",
        "v_zbierke",
        "zbierka",
        "owned",
        "rezervovane",
        "reserved",
    ),
    SOLD: ("predane", "predany", "predana", "sold"),
    WISH: ("chcem", "zelane", "zelany", "chcem_kupit", "wishlist", "wish", "want"),
}
_CONDITION = {
    ItemCondition.NEW_SEALED: (
        "nove_v_krabici",
        "nove",
        "novy",
        "nova",
        "v_krabici",
        "zatvorene",
        "nerozbalene",
        "new_sealed",
        "new",
        "sealed",
        "misb",
        "nisb",
    ),
    ItemCondition.OPENED_UNBUILT: (
        "rozbalene_nepostavene",
        "rozbalene",
        "otvorene",
        "otvoreny",
        "opened_unbuilt",
        "opened",
    ),
    ItemCondition.BUILT: ("postavene", "postaveny", "zlozene", "built", "used", "pouzite"),
    ItemCondition.PARTED_OUT: ("rozobrane", "rozobraty", "na_diely", "parted_out", "parted"),
}
_PURPOSE = {
    ItemPurpose.INVESTMENT: ("investicia", "investment"),
    ItemPurpose.FOR_SALE: ("na_predaj", "predaj", "for_sale"),
    ItemPurpose.DISPLAY: ("vystavene", "na_vystavenie", "display"),
    ItemPurpose.BUILD: ("na_stavanie", "stavanie", "build"),
}
_FLAGS = {
    "has_box": ("krabica", "s_krabicou", "box", "has_box"),
    "has_manual": ("manual", "navod", "has_manual", "instructions"),
    "has_stand": ("stojan", "has_stand", "stand"),
    "damaged_box": ("poskodena_krabica", "damaged_box"),
    "missing_parts": ("chybaju_dieliky", "chyba_dielik", "missing_parts"),
    "complete": ("kompletny", "kompletne", "kompletna", "complete"),
}
assert set(_FLAGS) == set(COLLECTION_FLAGS)

#: Slovenské názvy do šablóny a do chýb.
CONDITION_LABELS = {
    ItemCondition.NEW_SEALED: "nové v krabici",
    ItemCondition.OPENED_UNBUILT: "rozbalené",
    ItemCondition.BUILT: "postavené",
    ItemCondition.PARTED_OUT: "rozobrané",
}
PURPOSE_LABELS = {
    ItemPurpose.INVESTMENT: "investícia",
    ItemPurpose.FOR_SALE: "na predaj",
    ItemPurpose.DISPLAY: "vystavené",
    ItemPurpose.BUILD: "na stavanie",
}
OWNERSHIP_LABELS = {OWNED: "vlastnené", SOLD: "predané", WISH: "chcem"}
FLAG_LABELS = {
    "has_box": "krabica",
    "has_manual": "manuál",
    "has_stand": "stojan",
    "damaged_box": "poškodená krabica",
    "missing_parts": "chýbajú dieliky",
    "complete": "kompletný",
}


def _lookup(table: dict, raw: str):
    key = norm(raw)
    for value, aliases in table.items():
        if key in aliases or key == norm(str(value)):
            return value
    return None


# --- čítanie súboru ------------------------------------------------------------------


def read_cells(filename: str, content: bytes) -> list[list[Any]]:
    """Bunky prvého hárka (Excel) alebo riadky CSV."""
    if len(content) > MAX_BYTES:
        raise ImportFileError("Súbor má viac než 5 MB. Rozdeľ ho na menšie časti.")
    if not content.strip():
        raise ImportFileError("Súbor je prázdny.")
    lower = filename.lower()
    if content.startswith(b"PK\x03\x04") or lower.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(content)
    if content.startswith(b"\xd0\xcf\x11\xe0") or lower.endswith(".xls"):
        raise ImportFileError(
            "Starý formát Excelu (.xls) import nečíta. V Exceli daj Uložiť ako → .xlsx alebo CSV."
        )
    return _read_csv(content)


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # poškodený alebo iný súbor s príponou .xlsx
        raise ImportFileError("Súbor sa nedá otvoriť ako Excel (.xlsx).") from exc
    try:
        sheet = book.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def _decode(content: bytes) -> str:
    if content.startswith(b"\xef\xbb\xbf"):
        return content[3:].decode("utf-8", errors="replace")
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        # Slovenský Excel pri „Uložiť ako CSV“ píše vo Windows-1250.
        return content.decode("cp1250", errors="replace")


def _read_csv(content: bytes) -> list[list[Any]]:
    text = _decode(content)
    first = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(";,\t", key=first.count)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


# --- hlavička a riadky ------------------------------------------------------------


@dataclass(slots=True)
class ParsedRow:
    line: int
    values: dict[str, Any]
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass(slots=True)
class ParsedFile:
    rows: list[ParsedRow]
    ignored_columns: list[str]


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def parse(filename: str, content: bytes) -> ParsedFile:
    cells = read_cells(filename, content)
    header_at = next((i for i, row in enumerate(cells) if any(not _blank(c) for c in row)), None)
    if header_at is None:
        raise ImportFileError("V súbore nie sú žiadne riadky.")

    columns: dict[int, str] = {}
    ignored: list[str] = []
    for position, title in enumerate(cells[header_at]):
        if _blank(title):
            continue
        key = _COLUMN_OF.get(norm(title))
        if key is None or key in columns.values():
            ignored.append(str(title).strip())
        else:
            columns[position] = key
    if "catalog_num" not in columns.values():
        raise ImportFileError(
            "V prvom riadku chýba stĺpec s číslom setu. Pomenuj ho „cislo_setu“ "
            "(alebo „číslo“, „set“), prípadne použi šablónu."
        )

    rows: list[ParsedRow] = []
    for index, raw in enumerate(cells[header_at + 1 :], start=header_at + 2):
        if all(_blank(c) for c in raw):
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportFileError(f"Súbor má viac než {MAX_ROWS} riadkov. Rozdeľ ho.")
        values = {key: raw[pos] if pos < len(raw) else None for pos, key in columns.items()}
        rows.append(parse_row(index, values))
    if not rows:
        raise ImportFileError("Súbor má len hlavičku, žiadne riadky so setmi.")
    return ParsedFile(rows=rows, ignored_columns=ignored)


def parse_row(line: int, raw: dict[str, Any]) -> ParsedRow:
    """Jeden riadok na hodnoty pripravené na uloženie (reťazce, JSON)."""
    row = ParsedRow(line=line, values={})
    v = row.values

    def text(key: str, limit: int) -> str | None:
        value = raw.get(key)
        if _blank(value):
            return None
        return _cell_text(value)[:limit]

    def money(key: str, label: str, positive: bool = False) -> str | None:
        value = raw.get(key)
        if _blank(value):
            return None
        try:
            amount = parse_money(value)
        except ValueError:
            row.errors.append(f"{label}: „{_cell_text(value)}“ nie je suma.")
            return None
        if amount < 0 or (positive and amount == 0):
            row.errors.append(f"{label} musí byť kladná.")
            return None
        return str(amount)

    def day(key: str, label: str) -> str | None:
        value = raw.get(key)
        if _blank(value):
            return None
        try:
            parsed = parse_date(value)
        except ValueError:
            row.errors.append(f"{label}: „{_cell_text(value)}“ nie je dátum (napr. 12.3.2019).")
            return None
        if parsed > date.today():
            row.errors.append(f"{label} je v budúcnosti.")
            return None
        return parsed.isoformat()

    num = parse_catalog_num(raw.get("catalog_num"))
    v["raw_num"] = num
    if not num:
        row.errors.append("Chýba číslo setu.")
    v["name_hint"] = text("name", 200)
    v["note"] = text("note", 500)
    if v["note"] and v["note"].strip().lower() == EXAMPLE_NOTE:
        row.errors.append("Vzorový riadok zo šablóny, zmaž ho.")

    v["purchase_price"] = money("purchase_price", "Kúpna cena")
    _currency(row, raw)
    v["purchase_date"] = day("purchase_date", "Dátum kúpy")
    v["purchase_place"] = text("purchase_place", 160)
    v["location"] = text("location", 120)
    v["box"] = text("box", 40)
    v["sold_price"] = money("sold_price", "Predajná cena", positive=True)
    v["sold_date"] = day("sold_date", "Dátum predaja")
    v["sold_via"] = text("sold_via", 80)
    v["sold_fees"] = money("sold_fees", "Poplatky")
    v["sold_shipping"] = money("sold_shipping", "Poštovné")
    v["target_price"] = money("target_price", "Cieľová cena", positive=True)

    ownership = raw.get("ownership")
    if _blank(ownership):
        # Bez stĺpca vlastníctva prezradí predaj predajná cena alebo dátum.
        sold_hint = v["sold_price"] is not None or v["sold_date"] is not None
        v["ownership"] = SOLD if sold_hint else OWNED
    else:
        v["ownership"] = _lookup(_OWNERSHIP, _cell_text(ownership))
        if v["ownership"] is None:
            v["ownership"] = OWNED
            row.errors.append(
                f"Vlastníctvo „{_cell_text(ownership)}“ nepoznám. "
                "Použi vlastnené, predané alebo chcem."
            )

    condition = raw.get("condition")
    v["condition"] = str(ItemCondition.NEW_SEALED)
    if not _blank(condition):
        found = _lookup(_CONDITION, _cell_text(condition))
        if found is None:
            names = ", ".join(CONDITION_LABELS.values())
            row.errors.append(f"Stav „{_cell_text(condition)}“ nepoznám. Použi: {names}.")
        else:
            v["condition"] = str(found)

    purpose = raw.get("purpose")
    v["purpose"] = None
    if not _blank(purpose):
        found = _lookup(_PURPOSE, _cell_text(purpose))
        if found is None:
            names = ", ".join(PURPOSE_LABELS.values())
            row.errors.append(f"Zoznam „{_cell_text(purpose)}“ nepoznám. Použi: {names}.")
        else:
            v["purpose"] = str(found)

    flags: list[str] = []
    raw_flags = raw.get("flags")
    if not _blank(raw_flags):
        for part in re.split(r"[,;/|+]", _cell_text(raw_flags)):
            if not part.strip():
                continue
            found = _lookup(_FLAGS, part)
            if found is None:
                row.warnings.append(f"Príznak „{part.strip()}“ nepoznám, vynechám ho.")
            elif found not in flags:
                flags.append(found)
    v["flags"] = flags

    quantity = raw.get("quantity")
    v["quantity"] = 1
    if not _blank(quantity):
        try:
            v["quantity"] = parse_quantity(quantity)
        except ValueError:
            row.errors.append(f"Počet „{_cell_text(quantity)}“ musí byť celé číslo od 1 do 99.")

    _check_consistency(row)
    return row


#: Znak meny pri sume („1 290 Kč“), ktorý suma nepotrebuje.
_CURRENCY_MARK = re.compile(r"(?i)k[čc]|z[łl]|ft|chf|czk|usd|gbp|pln|huf|eur|[$£€]")


def _currency(row: ParsedRow, raw: dict[str, Any]) -> None:
    """Kúpa v cudzej mene: ``mena_kupy`` a ``kupna_cena_v_mene``.

    Eurá sa prepočítajú až pri potvrdení kurzom zo dňa kúpy, a len keď
    chýba ``kupna_cena_eur`` (export ju nesie, takže nahratý späť sa
    nemení). Euro v stĺpci meny je obyčajná kúpna cena v eurách.
    """
    from lego_api.services.currency import EUR, SUPPORTED, normalize

    v = row.values
    v["purchase_currency"] = None
    v["purchase_price_original"] = None
    code_raw = raw.get("purchase_currency")
    code: str | None = None
    if not _blank(code_raw):
        try:
            code = normalize(_cell_text(code_raw))
        except ValueError:
            row.errors.append(
                f"Menu „{_cell_text(code_raw)}“ nepoznám. Použi: {', '.join(SUPPORTED)}."
            )
            return
    amount_raw = raw.get("purchase_price_original")
    if _blank(amount_raw):
        return
    try:
        cleaned = _CURRENCY_MARK.sub("", amount_raw) if isinstance(amount_raw, str) else amount_raw
        amount = parse_money(cleaned)
    except ValueError:
        row.errors.append(f"Kúpna cena v mene: „{_cell_text(amount_raw)}“ nie je suma.")
        return
    if amount < 0:
        row.errors.append("Kúpna cena v mene musí byť kladná.")
        return
    if code is None:
        row.errors.append("Kúpna cena v mene potrebuje aj menu (stĺpec mena_kupy).")
        return
    if code == EUR:
        if v["purchase_price"] is None:
            v["purchase_price"] = str(amount)
        return
    v["purchase_currency"] = code
    v["purchase_price_original"] = str(amount)


def _check_consistency(row: ParsedRow) -> None:
    v = row.values
    if v["ownership"] == SOLD:
        if v["sold_price"] is None and not any("Predajná cena" in e for e in row.errors):
            row.errors.append("Predaný kus potrebuje predajnú cenu.")
        if v["sold_date"] is None and not any("Dátum predaja" in e for e in row.errors):
            row.errors.append("Predaný kus potrebuje dátum predaja, inak by v grafe visel navždy.")
        if v["sold_date"] and v["purchase_date"] and v["sold_date"] < v["purchase_date"]:
            row.errors.append("Dátum predaja je skôr než dátum kúpy.")
    elif v["sold_price"] is not None or v["sold_date"] is not None:
        row.warnings.append("Predajné údaje sa ignorujú, kus nie je označený ako predaný.")
    if v["ownership"] == WISH:
        if v["quantity"] != 1:
            row.warnings.append("V Chcem je set najviac raz, počet sa ignoruje.")
        bought = (
            v["purchase_price"]
            or v["purchase_price_original"]
            or v["purchase_date"]
            or v["purchase_place"]
        )
        if bought:
            row.warnings.append("Kúpne údaje sa pri Chcem ignorujú.")
    elif v["target_price"] is not None:
        row.warnings.append("Cieľová cena platí len pre Chcem, ignoruje sa.")


# --- hodnoty ----------------------------------------------------------------------


def _cell_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_catalog_num(value: Any) -> str:
    if _blank(value):
        return ""
    text = _cell_text(value).lower().replace(" ", "")
    return text if re.fullmatch(r"[0-9a-z][0-9a-z.\-]*", text) else ""


def parse_money(value: Any) -> Decimal:
    """39,99 / 39.99 / 1 234,50 / 1.234,50 / 39,99 € / bunka s číslom."""
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, int | float | Decimal):
        amount = Decimal(str(value))
    else:
        text = re.sub(r"(?i)eur|€|\s| ", "", str(value))
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
        try:
            amount = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError(value) from exc
    if not amount.is_finite():
        raise ValueError(value)
    return amount.quantize(Decimal("0.01"))


_DATE_FORMS = (
    (re.compile(r"^(\d{1,2})\.\s*(\d{1,2})\.\s*(\d{4})$"), ("d", "m", "y")),
    (re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?$"), ("y", "m", "d")),
    (re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$"), ("d", "m", "y")),
    (re.compile(r"^(\d{1,2})-(\d{1,2})-(\d{4})$"), ("d", "m", "y")),
)


def parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_text(value)
    for pattern, order in _DATE_FORMS:
        match = pattern.match(text)
        if match:
            parts = dict(zip(order, (int(g) for g in match.groups()), strict=True))
            return date(parts["y"], parts["m"], parts["d"])
    raise ValueError(value)


def parse_quantity(value: Any) -> int:
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(value)
        number = int(value)
    elif isinstance(value, int):
        number = value
    else:
        match = re.fullmatch(r"\s*(\d+)\s*(ks|kusy|kusov|x)?\.?\s*", str(value), re.IGNORECASE)
        if not match:
            raise ValueError(value)
        number = int(match.group(1))
    if not 1 <= number <= MAX_QUANTITY:
        raise ValueError(value)
    return number

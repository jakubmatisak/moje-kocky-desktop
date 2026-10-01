"""Šablóna na import v Exceli a v CSV.

Stĺpce sú tie isté, ktoré číta import a píše export. Vzorové riadky majú
poznámku ``EXAMPLE_NOTE``; keby ich niekto zabudol zmazať, import ich
ukáže ako chybu a preskočí, nič z nich nevznikne.
"""

from __future__ import annotations

import csv
import io
from datetime import date

from lego_api.services.currency import SUPPORTED
from lego_api.services.import_file import (
    CONDITION_LABELS,
    EXAMPLE_NOTE,
    FLAG_LABELS,
    OWNERSHIP_LABELS,
    PURPOSE_LABELS,
)

#: (stĺpec, popis do návodu, šírka v Exceli)
COLUMNS: list[tuple[str, str, int]] = [
    (
        "cislo_setu",
        "Katalógové číslo z krabice. Povinné. 10294 aj 10294-1; "
        "figúrka zo série s pomlčkou (71046-3).",
        12,
    ),
    ("pocet", "Koľko kusov vznikne. Prázdne = 1.", 7),
    (
        "vlastnictvo",
        "vlastnené, predané alebo chcem. Prázdne = vlastnené "
        "(pri vyplnenej predajnej cene predané).",
        13,
    ),
    ("stav", "nové v krabici, rozbalené, postavené, rozobrané. Prázdne = nové v krabici.", 16),
    ("zoznam", "investícia, na predaj, vystavené, na stavanie. Nepovinné.", 13),
    ("umiestnenie", "Kde kus leží, voľný text (Povala, Obývačka).", 14),
    ("krabica", "Číslo alebo názov krabice v miestnosti. Nepovinné.", 9),
    ("priznaky", "Oddelené čiarkou: " + ", ".join(FLAG_LABELS.values()) + ".", 20),
    ("kupna_cena_eur", "Cena za jeden kus, v eurách. 39,99 aj 39.99.", 14),
    (
        "mena_kupy",
        "Len pri kúpe v inej mene: CZK, USD, GBP, PLN, HUF alebo CHF. Prázdne = eurá.",
        10,
    ),
    (
        "kupna_cena_v_mene",
        "Len pri kúpe v inej mene: cena za kus v tej mene. Bez kúpnej ceny v eurách "
        "sa prepočíta kurzom ECB zo dňa kúpy.",
        16,
    ),
    ("datum_kupy", "Deň kúpy, 12.3.2019 alebo 2019-03-12.", 12),
    ("kde_kupene", "Obchod alebo predajca, voľný text.", 14),
    ("predajna_cena_eur", "Len pri predaných: za koľko sa predal.", 16),
    ("datum_predaja", "Len pri predaných: kedy. Povinné pri predanom.", 13),
    ("kanal_predaja", "Len pri predaných: Aukro, Bazoš, osobne…", 14),
    ("poplatky_eur", "Len pri predaných: poplatky trhoviska, ktoré si platil.", 12),
    ("postovne_eur", "Len pri predaných: poštovné, ktoré si platil.", 12),
    ("cielova_cena_eur", "Len pri Chcem: za koľko by si ho kúpil.", 15),
    ("poznamka", "Voľný text.", 30),
]
HEADER = [name for name, _, _ in COLUMNS]

_EXAMPLES: list[dict] = [
    {
        "cislo_setu": "10294",
        "pocet": 1,
        "vlastnictvo": "vlastnené",
        "stav": "nové v krabici",
        "zoznam": "investícia",
        "umiestnenie": "Povala",
        "krabica": "3",
        "priznaky": "krabica, manuál",
        "kupna_cena_eur": 599.99,
        "datum_kupy": date(2022, 3, 15),
        "kde_kupene": "LEGO Store",
    },
    {
        "cislo_setu": "75192",
        "pocet": 1,
        "vlastnictvo": "predané",
        "stav": "postavené",
        "priznaky": "krabica, manuál",
        "kupna_cena_eur": 649,
        "datum_kupy": date(2019, 10, 1),
        "kde_kupene": "Alza",
        "predajna_cena_eur": 820,
        "datum_predaja": date(2024, 2, 5),
        "kanal_predaja": "Bazoš",
        "poplatky_eur": 0,
        "postovne_eur": 9.9,
    },
    {"cislo_setu": "21318", "vlastnictvo": "chcem", "cielova_cena_eur": 180},
]


def _csv_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return f"{value.day}.{value.month}.{value.year}"
    if isinstance(value, float):
        return f"{value:.2f}".replace(".", ",")
    return str(value)


def template_csv() -> bytes:
    """Bodkočiarka a BOM, aby ho slovenský Excel otvoril s diakritikou."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(HEADER)
    for example in _EXAMPLES:
        row = {**example, "poznamka": EXAMPLE_NOTE}
        writer.writerow([_csv_value(row.get(name)) for name in HEADER])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Zbierka"
    sheet.append(HEADER)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="D01012")
    for position, (_, _, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=position)
        cell.font = bold
        cell.fill = fill
        sheet.column_dimensions[get_column_letter(position)].width = width
    sheet.freeze_panes = "A2"

    grey = Font(italic=True, color="808080")
    for example in _EXAMPLES:
        row = {**example, "poznamka": EXAMPLE_NOTE}
        sheet.append([row.get(name) for name in HEADER])
        for cell in sheet[sheet.max_row]:
            cell.font = grey

    #: Formát súm a dátumov na prvých riadkoch, výberové zoznamy na viac.
    formatted = 1000
    last = 5000
    column_of = {name: get_column_letter(i) for i, name in enumerate(HEADER, start=1)}
    for name in (
        "kupna_cena_eur",
        "kupna_cena_v_mene",
        "predajna_cena_eur",
        "poplatky_eur",
        "postovne_eur",
        "cielova_cena_eur",
    ):
        for row in sheet.iter_rows(
            min_row=2,
            max_row=formatted,
            min_col=HEADER.index(name) + 1,
            max_col=HEADER.index(name) + 1,
        ):
            row[0].number_format = "#,##0.00"
    for name in ("datum_kupy", "datum_predaja"):
        for row in sheet.iter_rows(
            min_row=2,
            max_row=formatted,
            min_col=HEADER.index(name) + 1,
            max_col=HEADER.index(name) + 1,
        ):
            row[0].number_format = "d.m.yyyy"

    # Výberové zoznamy: Excel ponúkne len hodnoty, ktoré import pozná.
    for name, labels in (
        ("vlastnictvo", OWNERSHIP_LABELS.values()),
        ("stav", CONDITION_LABELS.values()),
        ("zoznam", PURPOSE_LABELS.values()),
        ("mena_kupy", SUPPORTED),
    ):
        rule = DataValidation(type="list", formula1='"' + ",".join(labels) + '"', allow_blank=True)
        rule.error = "Vyber hodnotu zo zoznamu."
        rule.errorTitle = "Neznáma hodnota"
        sheet.add_data_validation(rule)
        rule.add(f"{column_of[name]}2:{column_of[name]}{last}")

    guide = book.create_sheet("Návod")
    guide.append(["Ako vyplniť šablónu"])
    guide["A1"].font = Font(bold=True, size=14)
    for line in (
        "Každý riadok je jeden set. Vyplň aspoň číslo setu, ostatné stĺpce môžu ostať prázdne.",
        "Sivé riadky sú príklady. Prepíš ich alebo zmaž; ak ostanú, import ich preskočí.",
        "Pred uložením sa ukáže náhľad: čo sa pridá, čo už v zbierke je a čo sa nedá prečítať.",
        "Import sa dá celý vrátiť v Nastaveniach → Import a export.",
        "Import nesťahuje ceny. Nové sety dostanú cenu tlačidlom obnovy cien v hornej lište.",
        "",
    ):
        guide.append([line])
    guide.append(["Stĺpec", "Čo doň patrí"])
    for cell in guide[guide.max_row]:
        cell.font = Font(bold=True)
    for name, description, _ in COLUMNS:
        guide.append([name, description])
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 100
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
